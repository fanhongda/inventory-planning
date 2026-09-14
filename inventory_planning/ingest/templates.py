"""
Templates — a blank workbook generated from a contract.

Some of what the planning needs has no source system to export it. Node topology,
which storage location is quarantine, a supersession list, the planner's own worksheet:
these are decisions, not records, and the measured scope of them is large — 46 of the
131 canonical fields are planner-local or derived and no ERP will ever hold them. Today
they are hand-edited JSON or they do not exist, and a person who will not open an editor
cannot supply any of them.

A template is the answer for exactly that data, and it is a small thing to build,
because **a template is an adapter whose headers we got to choose in advance.** The
sheet carries the canonical field names, so the mapping is right by construction: no
profiling, no scoring, no close call to adjudicate. It rides the ordinary routing path
with a hint rather than adding one of its own.

Three decisions worth stating.

**Generated from the contract, never hand-written.** A hand-maintained template is a
fourth place the field list lives — beside the contract, the adapter and the reader —
and it drifts on the first contract change, silently, in the direction of a column that
stops being filled. The same defect as a frozen adapter missing a new field.

**The workbook carries a fingerprint of the contract it came from.** A template filled
in last quarter against a contract that has since gained a required field should say so
on upload rather than arrive looking complete. The fingerprint covers what a person
filling the sheet would notice: field names, types, units, and which are required.

**A template is not a licence to re-key an export.** Where a real ERP export exists it
is the better input by a distance — it carries a header fingerprint, a null rate, a
contract test and a transform log, none of which survive a copy-paste. The dictionary
sheet says so, at the top, because the person most likely to reach for a template is
the person least able to see what re-keying costs.

Getting the mapping right is not the same as getting the numbers right. A template still
goes through contract tests, gates and the intake summary like anything else.
"""

from __future__ import annotations

import argparse
import hashlib
import sys
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any, Dict, List, Optional

from .contract import ContractRegistry, DocContract, default_registry

TEMPLATE_VERSION = 1

DATA_SHEET = "data"
DICTIONARY_SHEET = "dictionary"
META_SHEET = "_meta"

# Sheets the generator writes that are not data. Intake skips them by name rather than
# by shape: a two-column key/value sheet profiles perfectly well as a table, and a
# `_meta` sheet routed as a document would be a puzzling failure rather than an obvious
# one.
NON_DATA_SHEETS = (DICTIONARY_SHEET, META_SHEET)

_ADVICE = (
    "If this document has a real export from your ERP, upload that instead. An export "
    "carries where each column came from; a sheet filled in by hand does not, and the "
    "checks that catch a mis-read column have nothing to work with."
)


def contract_fingerprint(contract: DocContract) -> str:
    """
    A short hash of the parts of a contract a filled-in template depends on.

    Deliberately not the file's bytes: a reworded description or a new alias does not
    invalidate a sheet somebody has already filled in, and treating it as though it did
    would train people to ignore the warning. Names, types, units and requiredness are
    what change the work.
    """
    parts = []
    for name in sorted(contract.fields):
        spec = contract.fields[name]
        parts.append(f"{name}|{spec.type}|{spec.unit or ''}|{int(bool(spec.required))}")
    digest = hashlib.sha256("\n".join(parts).encode("utf-8")).hexdigest()
    return digest[:12]


def template_fields(contract: DocContract) -> List[Any]:
    """
    The contract's fields, required ones first.

    Required first because a template is read left to right and abandoned in the middle.
    Within each group the contract's own order is kept — it groups related fields, and
    re-sorting alphabetically would separate a quantity from its unit of measure.
    """
    ordered = list(contract.fields.values())
    return ([f for f in ordered if f.required]
            + [f for f in ordered if not f.required])


def _key_fields(contract: DocContract) -> set:
    return {str(k) for k in (contract.natural_key or [])}


def build_workbook(contract: DocContract, generated_on: date = None):
    """The three-sheet workbook for one contract. Returns an openpyxl Workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    generated_on = generated_on or date.today()
    fields = template_fields(contract)
    keys = _key_fields(contract)

    book = Workbook()
    data = book.active
    data.title = DATA_SHEET

    bold = Font(bold=True)
    required_fill = PatternFill("solid", fgColor="FFF2CC")
    for column, spec in enumerate(fields, start=1):
        cell = data.cell(row=1, column=column, value=spec.name)
        cell.font = bold
        if spec.required:
            cell.fill = required_fill
        data.column_dimensions[cell.column_letter].width = max(12, len(spec.name) + 4)
    data.freeze_panes = "A2"

    # The dictionary is as much the deliverable as the blank sheet is. Most of these
    # fields have a precise meaning that nothing currently writes down for the person
    # being asked to fill them in, and a column filled with the wrong thing is the
    # failure mode a template is supposed to remove rather than relocate.
    doc = book.create_sheet(DICTIONARY_SHEET)
    doc.cell(row=1, column=1, value=_ADVICE).font = Font(italic=True)
    doc.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    doc.cell(row=1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    doc.row_dimensions[1].height = 32

    headers = ["field", "required", "type", "unit", "part of the key",
               "what it means", "seen in exports as"]
    for column, title in enumerate(headers, start=1):
        doc.cell(row=3, column=column, value=title).font = bold
    for row, spec in enumerate(fields, start=4):
        doc.cell(row=row, column=1, value=spec.name)
        doc.cell(row=row, column=2, value="required" if spec.required else "")
        doc.cell(row=row, column=3, value=spec.type)
        doc.cell(row=row, column=4, value=spec.unit or "")
        doc.cell(row=row, column=5, value="key" if spec.name in keys else "")
        doc.cell(row=row, column=6, value=" ".join((spec.description or "").split()))
        doc.cell(row=row, column=7, value=", ".join(spec.aliases[:6]))
    for column, width in zip("ABCDEFG", (26, 10, 10, 22, 14, 70, 46)):
        doc.column_dimensions[column].width = width
    doc.freeze_panes = "A4"

    meta = book.create_sheet(META_SHEET)
    meta.cell(row=1, column=1, value="key").font = bold
    meta.cell(row=1, column=2, value="value").font = bold
    for row, (key, value) in enumerate(
        (
            ("template_version", TEMPLATE_VERSION),
            ("doc_type", contract.doc_type),
            ("contract_fingerprint", contract_fingerprint(contract)),
            ("generated_on", generated_on.isoformat()),
            ("fields", len(fields)),
            ("required_fields", sum(1 for f in fields if f.required)),
        ),
        start=2,
    ):
        meta.cell(row=row, column=1, value=key)
        meta.cell(row=row, column=2, value=value)
    meta.column_dimensions["A"].width = 24
    meta.column_dimensions["B"].width = 24
    return book


def emit(doc_type: str, out_dir, registry: ContractRegistry = None,
         generated_on: date = None) -> Path:
    """Write `template_<doc_type>.xlsx` into `out_dir`. Returns the path."""
    registry = registry or default_registry()
    contract = registry.get(doc_type)
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"template_{doc_type}.xlsx"
    build_workbook(contract, generated_on=generated_on).save(path)
    return path


@dataclass(frozen=True)
class TemplateMeta:
    """What a workbook says about itself, when it was generated by this module."""

    doc_type: str
    template_version: int = TEMPLATE_VERSION
    contract_fingerprint: str = ""
    generated_on: str = ""

    def staleness(self, registry: ContractRegistry = None) -> Optional[str]:
        """
        Why this template should be regenerated, or None when it is current.

        A stale template is not refused. It was filled in by a person and the data in it
        is real; what it cannot be is trusted to have had a column for everything the
        contract now asks for. So it loads, and it says what it is missing.
        """
        registry = registry or default_registry()
        try:
            contract = registry.get(self.doc_type)
        except KeyError:
            return (f"names doc_type {self.doc_type!r}, which no contract defines — "
                    f"generated by a different version of this package")
        if self.template_version != TEMPLATE_VERSION:
            return (f"was generated by template version {self.template_version}; "
                    f"this package writes version {TEMPLATE_VERSION}")
        current = contract_fingerprint(contract)
        if self.contract_fingerprint and self.contract_fingerprint != current:
            return (f"was generated from a different version of the {self.doc_type} "
                    f"contract ({self.contract_fingerprint} vs {current}) — a field may "
                    f"have been added since, and this sheet has no column for it")
        return None


def read_meta(path) -> Optional[TemplateMeta]:
    """
    The `_meta` sheet of a workbook, or None when there is not one.

    Never raises on a file that simply is not a template — every ordinary export takes
    this path on its way in, and a malformed workbook is the profiler's problem to
    report, not this module's to crash on.
    """
    path = Path(path)
    if path.suffix.lower() not in (".xlsx", ".xlsm", ".xls"):
        return None
    try:
        import pandas as pd

        book = pd.ExcelFile(path)
        if META_SHEET not in book.sheet_names:
            return None
        frame = book.parse(META_SHEET, dtype=str)
    except Exception:
        return None

    if frame.shape[1] < 2:
        return None
    pairs = {str(k).strip(): str(v).strip()
             for k, v in zip(frame.iloc[:, 0], frame.iloc[:, 1])
             if str(k).strip() and str(k).strip().lower() != "nan"}
    doc_type = pairs.get("doc_type", "")
    if not doc_type:
        return None
    try:
        version = int(float(pairs.get("template_version", TEMPLATE_VERSION)))
    except ValueError:
        version = -1
    return TemplateMeta(
        doc_type=doc_type,
        template_version=version,
        contract_fingerprint=pairs.get("contract_fingerprint", ""),
        generated_on=pairs.get("generated_on", ""),
    )


def main(argv: List[str] = None) -> int:
    parser = argparse.ArgumentParser(
        prog="python -m inventory_planning.ingest.templates",
        description="Generate a blank workbook from a document contract.",
    )
    parser.add_argument("--doc-type", action="append", default=[],
                        help="Contract to generate. Repeatable. Omit with --all.")
    parser.add_argument("--all", action="store_true",
                        help="Every contract this package defines")
    parser.add_argument("--output", default="templates", help="Directory to write into")
    parser.add_argument("--list", action="store_true", help="Show the contracts available")
    args = parser.parse_args(argv)

    registry = default_registry()
    if args.list:
        for doc_type in registry.doc_types:
            contract = registry.get(doc_type)
            required = sum(1 for f in contract.fields.values() if f.required)
            print(f"  {doc_type:<20} {len(contract.fields):>3} fields, "
                  f"{required} required   {contract_fingerprint(contract)}")
        return 0

    doc_types = registry.doc_types if args.all else list(args.doc_type)
    if not doc_types:
        parser.error("name at least one --doc-type, or pass --all (or --list)")

    for doc_type in doc_types:
        try:
            path = emit(doc_type, args.output, registry=registry)
        except KeyError as exc:
            print(f"  {exc}", file=sys.stderr)
            return 2
        print(f"  wrote {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
