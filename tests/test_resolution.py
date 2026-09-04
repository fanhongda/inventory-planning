"""
Resolution — that the structured view says the same thing the run does.

`explain.py` answers this question in prose for a person at a terminal. This answers it
in a shape an interface can lay out, and the risk in having two of them is that they
drift. It is contained by not re-implementing anything: the module calls the same
`Intake.load_frame` a real run calls, and reports what came back.
"""

import pandas as pd
import pytest

from inventory_planning.ingest.templates import emit
from inventory_planning.resolution import (
    SOURCE_ABSENT, SOURCE_DECLARED, SOURCE_DERIVED, SOURCE_MAPPED,
    resolve_file, resolve_frame,
)
from inventory_planning.store.declarations import Declarations, Override, SCOPE_MAPPING


@pytest.fixture
def inventory():
    return resolve_file("sample_data/inventory.csv")[0]


class TestWhatWasMadeOfAFile:

    def test_it_routes_and_counts(self, inventory):
        assert inventory.doc_type == "inventory"
        assert inventory.rows == 10
        assert inventory.tests_passed
        assert inventory.missing_required == []

    def test_each_field_says_where_it_came_from(self, inventory):
        by_field = {f.field: f for f in inventory.fields}
        assert by_field["sku"].source == SOURCE_MAPPED
        assert by_field["sku"].column == "Item Code"
        assert by_field["sku"].fill_rate == 1.0
        # Computed by the contract from other fields — not a column anybody supplied.
        assert by_field["qty_available"].source == SOURCE_DERIVED
        # The case the exposure ledger exists for: no column, and no complaint.
        assert by_field["currency"].source == SOURCE_ABSENT

    def test_required_fields_are_listed_first(self, inventory):
        required = [f.required for f in inventory.fields]
        assert required[0] is True
        assert required == sorted(required, reverse=True)

    def test_a_column_matching_no_field_is_named(self):
        frame = pd.read_csv("sample_data/inventory.csv", dtype=str)
        frame["Warehouse Notes"] = "n/a"
        resolved = resolve_frame(frame, source_name="inventory.csv")
        assert "Warehouse Notes" in resolved.unmatched_columns

    def test_a_missing_required_field_is_reported_not_raised(self):
        frame = pd.DataFrame({"On Hand": ["5", "6"], "Report Date": ["2026-01-01"] * 2})
        resolved = resolve_frame(frame, source_name="headless.csv",
                                 doc_type_hint="inventory")
        assert "sku" in [f.field for f in resolved.missing_required]


class TestDeclarationsAreVisibleAsSuch:

    def _declared(self, column):
        return Declarations(overrides=[Override(
            scope=SCOPE_MAPPING, field="sku", value=column,
            target={"doc_type": "inventory"}, reason="worked example", by="tests")])

    def test_a_field_taken_from_a_declared_column_says_declared(self):
        frame = pd.read_csv("sample_data/inventory.csv", dtype=str)
        resolved = resolve_frame(frame, source_name="inventory.csv",
                                 declarations=self._declared("Base Unit"))
        sku = {f.field: f for f in resolved.fields}["sku"]
        assert sku.source == SOURCE_DECLARED
        assert sku.column == "Base Unit"

    def test_a_declaration_naming_a_column_the_file_lacks_is_reported_unapplied(self):
        """
        The failure this guards: the routed column labelled `declared`, and silence.
        Someone then believes the mapping is corrected and stops looking.
        """
        frame = pd.read_csv("sample_data/inventory.csv", dtype=str)
        resolved = resolve_frame(frame, source_name="inventory.csv",
                                 declarations=self._declared("Description"))
        sku = {f.field: f for f in resolved.fields}["sku"]
        assert sku.source == SOURCE_MAPPED
        assert sku.column == "Item Code"
        assert resolved.ignored_declarations == [
            {"field": "sku", "column": "Description",
             "reason": "the file has no such column"}]


class TestAGeneratedTemplateNeedsNoInference:

    def test_a_filled_template_routes_on_its_own_meta_sheet(self, tmp_path):
        from openpyxl import load_workbook

        blank = emit("substitution", tmp_path)
        book = load_workbook(blank)
        for column, value in enumerate(
                ["P-1", "P-2", "supersede", 1, "2026-01-15", "renumbered"], start=1):
            book["data"].cell(row=2, column=column, value=value)
        book.save(blank)

        resolved = resolve_file(blank)
        assert len(resolved) == 1          # the dictionary and _meta sheets are skipped
        assert resolved[0].doc_type == "substitution"
        assert resolved[0].confidence == pytest.approx(1.0)

    def test_a_template_with_no_rows_in_it_has_nothing_to_resolve(self, tmp_path):
        """
        Headers alone are not data. `/uploads` turns this into a sentence naming the
        template and what to do; here it is simply an empty list.
        """
        assert resolve_file(emit("substitution", tmp_path)) == []


class TestItSerialisesWholesale:

    def test_to_dict_carries_the_fields_an_interface_lays_out(self, inventory):
        body = inventory.to_dict()
        assert body["doc_type"] == "inventory"
        assert {"field", "source", "column", "fill_rate", "required"} <= set(
            body["fields"][0])
        for key in ("missing_required", "unmatched_columns", "ignored_declarations",
                    "transforms", "tests_passed"):
            assert key in body
