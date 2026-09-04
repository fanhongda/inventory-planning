"""
Templates — that the blank sheet is the contract, and knows when it stopped being it.

A template is an adapter whose headers were chosen in advance, so the two failures worth
testing are the two that would make it stop being one: headers drifting away from the
canonical field names, and a sheet filled in against a contract that has since changed
arriving as though it were complete.
"""

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from inventory_planning.ingest.contract import DocContract, default_registry
from inventory_planning.ingest.intake import Intake
from inventory_planning.ingest.templates import (
    DATA_SHEET, DICTIONARY_SHEET, META_SHEET, NON_DATA_SHEETS, TEMPLATE_VERSION,
    build_workbook, contract_fingerprint, emit, read_meta, template_fields,
)

REGISTRY = default_registry()
DOC_TYPES = REGISTRY.doc_types


def _contract(**field_specs) -> DocContract:
    return DocContract.parse({
        "doc_type": "trial",
        "fields": field_specs,
    })


@pytest.fixture
def substitution(tmp_path) -> Path:
    return emit("substitution", tmp_path, registry=REGISTRY)


class TestEveryShippedContractGenerates:

    @pytest.mark.parametrize("doc_type", DOC_TYPES)
    def test_three_sheets_and_the_canonical_headers(self, doc_type, tmp_path):
        """
        Parametrised over the live registry rather than a fixed list, so a contract
        added later is covered without anyone remembering to add it here.
        """
        path = emit(doc_type, tmp_path, registry=REGISTRY)
        book = load_workbook(path)
        assert book.sheetnames == [DATA_SHEET, DICTIONARY_SHEET, META_SHEET]

        contract = REGISTRY.get(doc_type)
        headers = [c.value for c in book[DATA_SHEET][1]]
        assert set(headers) == set(contract.fields)

    @pytest.mark.parametrize("doc_type", DOC_TYPES)
    def test_required_columns_come_first(self, doc_type, tmp_path):
        """A template is read left to right and abandoned in the middle."""
        contract = REGISTRY.get(doc_type)
        names = [f.name for f in template_fields(contract)]
        required = [n for n in names if contract.fields[n].required]
        assert names[:len(required)] == required


class TestTheDictionaryIsPartOfTheDeliverable:

    def test_it_carries_meaning_key_membership_and_the_headers_seen_in_the_wild(
            self, substitution):
        doc = pd.read_excel(substitution, sheet_name=DICTIONARY_SHEET, header=2)
        row = doc[doc["field"] == "old_sku"].iloc[0]
        assert row["required"] == "required"
        assert row["part of the key"] == "key"
        assert "retired" in str(row["what it means"])
        assert "old material" in str(row["seen in exports as"])

    def test_it_says_an_export_beats_a_re_keyed_sheet(self, substitution):
        """
        The person most likely to reach for a template is the one least able to see
        what re-keying an export costs, so the advice sits above the field list.
        """
        first = load_workbook(substitution)[DICTIONARY_SHEET].cell(row=1, column=1).value
        assert "upload that instead" in first


class TestTheFingerprintTracksWhatChangesTheWork:

    def test_regenerating_the_same_contract_gives_the_same_fingerprint(self):
        contract = REGISTRY.get("inventory")
        assert contract_fingerprint(contract) == contract_fingerprint(contract)

    def test_a_new_required_field_changes_it(self):
        before = _contract(sku={"type": "string", "required": True})
        after = _contract(sku={"type": "string", "required": True},
                          plant={"type": "string", "required": True})
        assert contract_fingerprint(before) != contract_fingerprint(after)

    def test_making_a_field_required_changes_it(self):
        before = _contract(sku={"type": "string"})
        after = _contract(sku={"type": "string", "required": True})
        assert contract_fingerprint(before) != contract_fingerprint(after)

    def test_a_reworded_description_or_a_new_alias_does_not(self):
        """
        A fingerprint that moves on every edit trains people to ignore the warning it
        raises. It covers what changes the filling-in, not what changes the reading.
        """
        before = _contract(sku={"type": "string", "description": "the material",
                                "aliases": ["material"]})
        after = _contract(sku={"type": "string", "description": "the material number",
                               "aliases": ["material", "matnr", "part no"]})
        assert contract_fingerprint(before) == contract_fingerprint(after)


class TestReadingBackWhatTheWorkbookSaysAboutItself:

    def test_a_generated_template_names_its_doc_type(self, substitution):
        meta = read_meta(substitution)
        assert meta.doc_type == "substitution"
        assert meta.template_version == TEMPLATE_VERSION
        assert meta.staleness(REGISTRY) is None

    def test_an_ordinary_export_is_not_a_template_and_does_not_raise(self, tmp_path):
        """Every real export takes this path on its way in."""
        assert read_meta(Path("sample_data/inventory.csv")) is None
        plain = tmp_path / "export.xlsx"
        pd.DataFrame({"Material": ["A"], "Qty": [1]}).to_excel(plain, index=False)
        assert read_meta(plain) is None

    def test_a_workbook_from_an_older_contract_says_so(self, substitution, tmp_path):
        stale = tmp_path / "stale.xlsx"
        book = load_workbook(substitution)
        for row in book[META_SHEET].iter_rows(min_row=2):
            if row[0].value == "contract_fingerprint":
                row[1].value = "0000deadbeef"
        book.save(stale)
        message = read_meta(stale).staleness(REGISTRY)
        assert "different version of the substitution contract" in message

    def test_a_doc_type_no_contract_defines_says_so(self, substitution, tmp_path):
        unknown = tmp_path / "unknown.xlsx"
        book = load_workbook(substitution)
        for row in book[META_SHEET].iter_rows(min_row=2):
            if row[0].value == "doc_type":
                row[1].value = "warehouse_layout"
        book.save(unknown)
        assert "no contract defines" in read_meta(unknown).staleness(REGISTRY)


class TestAFilledTemplateGoesStraightIn:

    def _fill(self, blank: Path, target: Path, rows) -> Path:
        book = load_workbook(blank)
        sheet = book[DATA_SHEET]
        for r, row in enumerate(rows, start=2):
            for c, value in enumerate(row, start=1):
                sheet.cell(row=r, column=c, value=value)
        book.save(target)
        return target

    def test_it_routes_on_its_own_word_with_nothing_inferred(self, substitution, tmp_path):
        filled = self._fill(substitution, tmp_path / "filled.xlsx", [
            ["P-1001", "P-2001", "supersede", 1, "2026-01-15", "renumbered"],
            ["P-1002", "P-2002", "phase", 1, "2026-03-01", "ramping"],
        ])
        result = Intake(verbose=False).load_files([filled])
        assert list(result.documents) == ["substitution"]
        doc = result.documents["substitution"]
        assert doc.route.confidence == pytest.approx(1.0)
        assert doc.row_count == 2
        assert doc.passed
        assert result.failures == []

    def test_the_companion_sheets_are_not_treated_as_documents(self, substitution,
                                                               tmp_path):
        """
        A two-column key/value sheet profiles as a perfectly good table. Skipped by
        name, so `_meta` never becomes a document that fails somewhere confusing.
        """
        filled = self._fill(substitution, tmp_path / "filled.xlsx",
                            [["P-1", "P-2", "supersede", 1, "2026-01-15", "x"]])
        result = Intake(verbose=False).load_files([filled])
        assert all(name not in str(result.failures) for name in NON_DATA_SHEETS)
        assert len(result.documents) == 1

    def test_a_stale_template_is_loaded_and_reported_not_refused(self, substitution,
                                                                 tmp_path):
        """
        The rows were filled in by a person and are real. What they cannot be trusted to
        have is a column for everything the contract asks for now.
        """
        filled = self._fill(substitution, tmp_path / "filled.xlsx",
                            [["P-1", "P-2", "supersede", 1, "2026-01-15", "x"]])
        book = load_workbook(filled)
        for row in book[META_SHEET].iter_rows(min_row=2):
            if row[0].value == "contract_fingerprint":
                row[1].value = "0000deadbeef"
        book.save(filled)

        result = Intake(verbose=False).load_files([filled])
        assert result.documents["substitution"].row_count == 1
        assert any("different version" in n for n in result.notes)


class TestBuildWorkbookNeedsNoDisk:

    def test_it_returns_a_workbook(self):
        book = build_workbook(REGISTRY.get("inventory"))
        assert book.sheetnames == [DATA_SHEET, DICTIONARY_SHEET, META_SHEET]
