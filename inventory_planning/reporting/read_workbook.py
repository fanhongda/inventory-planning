"""
Reading a run's workbook back, so a screen can show it without rebuilding it.

INTERFACE.md §7 revised the position that results get no screen. The objection was that
a screen showing the same numbers becomes a second place they are formatted, rounded and
subtly disagreed about — which holds only if the screen *re-derives* them. It does not
if the screen renders the run's own workbook. One artefact, two renderings.

So the constraint this module exists to hold: **it reads the file and computes nothing.**
No sums, no ratios, no re-ordering of columns, no filling of blanks. A cell's value is
whatever `workbook.py` wrote into it.

The one thing that is not simply copied is how a number is *displayed*, and that is the
subtle half of the objection rather than an exception to it. `workbook.py` writes a raw
value and a `number_format`, and Excel shows the two combined: `12345.67` under `#,##0`
reads as `12,346`. A screen that printed the raw value would disagree with the workbook
on every money column while holding the identical number — the disagreement §7 warns
about, arrived at from the other direction. So the format is read off the cell and
applied, and the formats understood here are exactly the ones the writer emits.

`test_results_screen.py` pins that correspondence: every format constant in `workbook.py`
must be one this module renders. Add a fifth format there and the test fails here, which
is the only way this stays true without the two files being one file.

Rounding is half away from zero, because that is what Excel does. Python's `round` is
half to even, so `round(0.5)` is `0` — and a screen that showed `0` where the workbook
shows `1` would be the disagreement in its purest form.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field as dc_field
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Dict, List, Optional

# The shapes `workbook.py` emits: an optional thousands group, a required integer part,
# an optional decimal part, an optional trailing percent. Parsed rather than looked up
# by exact string so a near variant (`#,##0.00`) renders correctly instead of falling
# back to raw — but still narrow enough that anything unexpected is visibly not handled.
_FORMAT = re.compile(r"^(?P<group>#,##)?0(?:\.(?P<decimals>0+))?(?P<percent>%)?$")


class WorkbookUnreadable(ValueError):
    """The file named by the manifest is not there, or is not a workbook."""


@dataclass
class Sheet:
    """One sheet, as far as it was read."""

    name: str
    columns: List[str] = dc_field(default_factory=list)
    rows: List[List[Any]] = dc_field(default_factory=list)
    total: int = 0
    offset: int = 0
    # Format strings this module did not understand, by column. Reported rather than
    # swallowed: a column shown raw where the workbook shows it formatted is a visible
    # difference, and the screen should be able to say so instead of looking wrong.
    unformatted: List[str] = dc_field(default_factory=list)

    def to_dict(self) -> Dict[str, Any]:
        return {"name": self.name, "columns": self.columns, "rows": self.rows,
                "total": self.total, "offset": self.offset,
                "unformatted": self.unformatted}


def render_number(value: Any, number_format: str) -> Optional[str]:
    """
    One number as the workbook displays it, or None if the format is not understood.

    None rather than a guess: a wrong rendering is worse than an unrendered one, because
    the first is believed.
    """
    match = _FORMAT.match((number_format or "").strip())
    if match is None:
        return None
    try:
        amount = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        return None

    if match.group("percent"):
        amount *= 100
    decimals = len(match.group("decimals") or "")
    quantum = Decimal(1).scaleb(-decimals)
    # Half away from zero, which is Excel's rule. `round()` is half to even and would
    # show 0 where the workbook shows 1.
    amount = amount.quantize(quantum, rounding=ROUND_HALF_UP)

    text = f"{amount:,f}" if match.group("group") else f"{amount:f}"
    if decimals == 0 and "." in text:
        text = text.split(".")[0]
    return text + ("%" if match.group("percent") else "")


def _cell_value(cell) -> Any:
    """
    The cell as JSON, formatted by its own `number_format` where that is understood.

    A formula is reported as the formula text rather than evaluated. `workbook.py`
    writes no formulas, so this never fires on a workbook this pipeline produced — but
    evaluating one would be the screen computing, which is the line this module exists
    not to cross.
    """
    value = cell.value
    if value is None:
        return None
    if isinstance(value, str):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float, Decimal)):
        rendered = render_number(value, cell.number_format)
        return rendered if rendered is not None else value
    return str(value)


def sheets(path) -> List[Dict[str, Any]]:
    """Every sheet in the workbook with its size, without reading the cells."""
    book = _open(path)
    try:
        out = []
        for name in book.sheetnames:
            worksheet = book[name]
            out.append({
                "name": name,
                # `max_row` counts the header, which is not a row of results.
                "rows": max(0, (worksheet.max_row or 0) - 1),
                "columns": worksheet.max_column or 0,
            })
        return out
    finally:
        book.close()


def read_sheet(path, sheet: str, offset: int = 0, limit: int = 200) -> Sheet:
    """
    One window of one sheet, values as the workbook holds and displays them.

    Paged because these sheets run to thousands of rows and a screen that fetched all of
    them would be slower than opening the file in Excel, which is the thing it is meant
    to save.
    """
    book = _open(path)
    try:
        if sheet not in book.sheetnames:
            raise WorkbookUnreadable(
                f"{Path(path).name} has no sheet {sheet!r}. It has: "
                f"{', '.join(book.sheetnames)}.")
        worksheet = book[sheet]
        rows = worksheet.iter_rows()
        try:
            header = [str(c.value) if c.value is not None else ""
                      for c in next(rows)]
        except StopIteration:
            return Sheet(name=sheet)

        total = max(0, (worksheet.max_row or 0) - 1)
        offset = max(0, int(offset))
        window: List[List[Any]] = []
        unformatted: Dict[str, None] = {}
        for index, row in enumerate(rows):
            if index < offset:
                continue
            if len(window) >= limit:
                break
            values = []
            for position, cell in enumerate(row):
                values.append(_cell_value(cell))
                if (isinstance(cell.value, (int, float, Decimal))
                        and not isinstance(cell.value, bool)
                        and (cell.number_format or "General") != "General"
                        and render_number(cell.value, cell.number_format) is None
                        and position < len(header)):
                    unformatted[header[position]] = None
            window.append(values)

        # `max_row` comes from the sheet's declared dimension, and a workbook written
        # without one reports None in read-only mode. openpyxl writes it and these are
        # its workbooks, so this never fires here — but "rows 1–10 of 0" is the kind of
        # nonsense a screen should not be able to print.
        total = max(total, offset + len(window))
        return Sheet(name=sheet, columns=header, rows=window, total=total,
                     offset=offset, unformatted=list(unformatted))
    finally:
        book.close()


def _open(path):
    """
    Read-only, values and formats, no formula evaluation.

    `read_only` matters at this size — the planning sheet is wide and openpyxl's normal
    mode builds every cell object up front — and `data_only=False` is deliberate: a
    cached formula result would be a number this pipeline never wrote.
    """
    path = Path(path)
    if not path.exists():
        raise WorkbookUnreadable(f"{path} is not there any more.")
    try:
        import openpyxl

        return openpyxl.load_workbook(path, read_only=True, data_only=False)
    except WorkbookUnreadable:
        raise
    except Exception as exc:                      # noqa: BLE001 - reported, not handled
        raise WorkbookUnreadable(f"{path.name} could not be opened: {exc}") from exc
